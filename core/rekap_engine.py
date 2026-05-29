"""
Engine untuk Modul ④ Rekap UP3 - Generate summary report per owner.
Optimized dengan openpyxl untuk formatting dan formula.
"""

import os
from pathlib import Path
from typing import Dict, List, Optional, Any, Tuple, Set
import pandas as pd
from collections import defaultdict
from copy import copy as _copy

from core.base_engine import BaseEngine
from utils import clean_filename, find_column_index, safe_float
from config.constants import COLUMNS


def _load_rekap_config(template_path: str) -> Tuple[Dict, Dict, List, str]:
    """
    Load CONFIG sheet + UNIT_MASTER dari template file.
    
    Returns: (config_map, owner_map, template_rows, template_name)
    """
    from openpyxl import load_workbook
    
    config_map = {}
    owner_map = {}
    template_rows = []
    
    wb = load_workbook(template_path, data_only=True)
    
    # --- UNIT_MASTER ---
    if "UNIT_MASTER" in wb.sheetnames:
        df_owner = pd.read_excel(template_path, sheet_name="UNIT_MASTER", dtype=str)
        df_owner.columns = [str(c).strip() for c in df_owner.columns]
        
        for _, row in df_owner.iterrows():
            code = str(row.get("OWNER_CODE", "")).strip()
            name = str(row.get("NAME_UP3", "")).strip()
            
            if code and code != "nan" and name and name != "nan":
                # Add multiple key variants
                for key in [code.upper(),
                           code.replace(" ", "").upper(),
                           code.replace(".", "").upper()]:
                    owner_map[key] = name
                
                # Add prefix variants
                for prefix in ["ULP.", "UP3."]:
                    if code.upper().startswith(prefix):
                        owner_map[code[len(prefix):].upper()] = name
    
    # --- CONFIG ---
    if "CONFIG" in wb.sheetnames:
        df_cfg = pd.read_excel(template_path, sheet_name="CONFIG", dtype=str)
        df_cfg.columns = [str(c).strip() for c in df_cfg.columns]
        H = [c.upper() for c in df_cfg.columns]
        
        def gc(cands): 
            return find_column_index(H, cands)
        
        iSub = gc(["SUBCLASS"])
        iKode = gc(["KODE_ASSET", "KODE ASSET"])
        iHit = gc(["JENIS HITUNGAN", "JENIS_HITUNGAN"])
        iKol = gc(["KOLOM"])
        iJns = gc(["JENIS ASSET", "JENIS_ASSET"])
        iKlp = gc(["KELOMPOK ASSET", "KELOMPOK_ASSET"])
        
        for _, row in df_cfg.iterrows():
            sub = str(row.iloc[iSub] if iSub is not None else "").strip().upper()
            kode = str(row.iloc[iKode] if iKode is not None else "").strip().upper()
            
            if sub and sub != "NAN" and kode and kode != "NAN":
                config_map[sub] = {
                    "kode_asset": kode,
                    "jenis_hitungan": str(row.iloc[iHit] if iHit is not None else "").strip().lower(),
                    "kolom_nilai": str(row.iloc[iKol] if iKol is not None else "").strip().upper(),
                    "jenis_asset": str(row.iloc[iJns] if iJns is not None else "").strip(),
                    "kelompok_asset": str(row.iloc[iKlp] if iKlp is not None else "").strip(),
                }
    
    # --- Template sheet {{UP3}} ---
    tmpl_name = "{{UP3}}"
    if tmpl_name not in wb.sheetnames:
        tmpl_name = next(
            (s for s in wb.sheetnames if s not in ["UNIT_MASTER", "CONFIG"]),
            None
        )
    
    if tmpl_name:
        ws = wb[tmpl_name]
        for r in range(9, ws.max_row + 1):
            jenis = ws.cell(r, 3).value
            if jenis and str(jenis).strip() not in ("", "None"):
                template_rows.append({
                    "row": r,
                    "no": str(ws.cell(r, 1).value or "").strip(),
                    "kelompok_asset": str(ws.cell(r, 2).value or "").strip(),
                    "jenis_asset": str(jenis).strip(),
                    "kode_asset": str(ws.cell(r, 4).value or "").strip().upper(),
                })
    
    wb.close()
    return config_map, owner_map, template_rows, tmpl_name


def _process_rekap_data(
    df: pd.DataFrame,
    config_map: Dict,
    owner_map: Dict
) -> Tuple[Dict, Set]:
    """
    Build summary: {owner: {kode_asset: {total, operating, inactive, itw, kms, kva, ...}}}
    """
    H = [str(c).upper().strip() for c in df.columns]
    
    iSub = find_column_index(H, ["SUBCLASS", "SUBCLASS_CODE", "SUBCLASS CODE"])
    iOwner = find_column_index(H, ["OWNER_ASSET", "OWNER_ASET", "OWNER"])
    iStatus = find_column_index(H, ["STATUS"])
    iEM = find_column_index(H, ["END_MEASURE", "PANJANG", "LENGTH"])
    iKap = find_column_index(H, ["KAPASITAS", "KAPASITAS_TRAFO", "RATED_CAPACITY", "CAPACITY", "KVA"])
    
    if iSub is None:
        raise ValueError("Kolom SUBCLASS tidak ditemukan di file master!")
    
    summary = {}
    not_found_subs = set()
    
    for _, row in df.iterrows():
        sub = str(row.iloc[iSub] or "").strip().upper()
        if not sub or sub == "NAN":
            continue
        
        if sub not in config_map:
            not_found_subs.add(sub)
            continue
        
        cfg = config_map[sub]
        kode = cfg["kode_asset"]
        
        # Get owner
        owner_raw = str(row.iloc[iOwner] or "").strip() if iOwner is not None else ""
        owner = owner_map.get(
            owner_raw.upper(),
            owner_map.get(
                owner_raw.replace(" ", "").upper(),
                owner_map.get(
                    owner_raw.replace(".", "").upper(),
                    owner_raw
                )
            )
        ) or "UNKNOWN"
        
        # Get status
        st = str(row.iloc[iStatus] or "").upper().strip() if iStatus is not None else ""
        if any(x in st for x in ("OPERATING", "OPERATE", "OPERASI")):
            sk = "operating"
        elif any(x in st for x in ("WAREHOUSE", "ITW", "GUDANG")):
            sk = "itw"
        elif any(x in st for x in ("INACTIVE", "NONAKTIF", "TIDAK AKTIF")):
            sk = "inactive"
        else:
            sk = "operating"
        
        # Get nilai
        nilai = 0.0
        kv = cfg["kolom_nilai"]
        if kv:
            ci = find_column_index(H, [kv])
            if ci is not None:
                nilai = safe_float(row.iloc[ci])
        elif "panjang" in cfg["jenis_hitungan"] and iEM is not None:
            nilai = safe_float(row.iloc[iEM])
        elif "kapasitas" in cfg["jenis_hitungan"] and iKap is not None:
            nilai = safe_float(row.iloc[iKap])
        
        summary.setdefault(owner, {})
        if kode not in summary[owner]:
            summary[owner][kode] = {
                "total": 0, "operating": 0, "inactive": 0, "itw": 0,
                "kms": 0.0, "kva": 0.0,
                "jenis_asset": cfg["jenis_asset"],
                "kelompok_asset": cfg["kelompok_asset"],
            }
        
        rec = summary[owner][kode]
        rec["total"] += 1
        rec[sk] += 1
        if "panjang" in cfg["jenis_hitungan"] and nilai > 0:
            rec["kms"] += nilai
        elif "kapasitas" in cfg["jenis_hitungan"] and nilai > 0:
            rec["kva"] += nilai
    
    return summary, not_found_subs


def _write_rekap_sheet(
    wb_out,
    ws_tpl,
    owner: str,
    uid: str,
    owner_data: Dict,
    template_rows: List,
    filter_mode: int
) -> str:
    """
    Tulis satu sheet rekap ke wb_out.
    """
    from openpyxl.utils import get_column_letter
    
    safe = clean_filename(owner, 31)
    ws = wb_out.create_sheet(title=safe[:31])
    
    # Copy template
    for r in range(1, ws_tpl.max_row + 1):
        for c in range(1, ws_tpl.max_column + 1):
            src = ws_tpl.cell(r, c)
            dst = ws.cell(r, c)
            dst.value = src.value
            if src.has_style:
                dst.font = _copy(src.font)
                dst.fill = _copy(src.fill)
                dst.border = _copy(src.border)
                dst.alignment = _copy(src.alignment)
                dst.number_format = src.number_format
    
    for mr in ws_tpl.merged_cells.ranges:
        ws.merge_cells(str(mr))
    
    for ci, cd in ws_tpl.column_dimensions.items():
        ws.column_dimensions[ci].width = cd.width
    
    for ri, rd in ws_tpl.row_dimensions.items():
        ws.row_dimensions[ri].height = rd.height
    
    # Replace header placeholders
    for r in range(1, 9):
        for c in range(1, ws_tpl.max_column + 1):
            cell = ws.cell(r, c)
            if isinstance(cell.value, str):
                cell.value = cell.value.replace("{{UP3}}", owner).replace("{{UID}}", uid or "")
    
    # Fill data
    rows_with_data = []
    rows_to_remove = []
    
    for tmpl in template_rows:
        rn = tmpl["row"]
        kode = tmpl["kode_asset"]
        rec = owner_data.get(kode) if kode else None
        has = bool(rec and rec["total"] > 0)
        
        if has:
            rows_with_data.append(rn)
            ws.cell(rn, 5).value = rec["total"] if rec["total"] > 0 else None
            ws.cell(rn, 6).value = rec["operating"] if rec["operating"] > 0 else None
            ws.cell(rn, 7).value = rec["inactive"] if rec["inactive"] > 0 else None
            ws.cell(rn, 8).value = rec["itw"] if rec["itw"] > 0 else None
            ws.cell(rn, 9).value = round(rec["kms"], 3) if rec["kms"] > 0 else None
            ws.cell(rn, 10).value = round(rec["kva"], 3) if rec["kva"] > 0 else None
            
            fmt = "#,##0"
            for ci in range(5, 11):
                ws.cell(rn, ci).number_format = fmt
        else:
            if filter_mode == 1:
                rows_to_remove.append(rn)
    
    # Remove empty rows (mode 1)
    if filter_mode == 1 and rows_to_remove:
        mrs_to_remove = []
        for mr in list(ws.merged_cells.ranges):
            if any(mr.min_row <= r <= mr.max_row for r in rows_to_remove):
                mrs_to_remove.append(str(mr))
        for mr_str in mrs_to_remove:
            try:
                ws.unmerge_cells(mr_str)
            except Exception:
                pass
        for r in sorted(rows_to_remove, reverse=True):
            ws.delete_rows(r)
    
    # Renumber
    data_start = 9
    last_row = ws.max_row
    no_counter = 1
    for r in range(data_start, last_row):
        cell_a = ws.cell(r, 1)
        cell_c = ws.cell(r, 3)
        if cell_c.value and str(cell_c.value).strip() not in ("", "None"):
            if cell_a.value is not None and str(cell_a.value).strip().isdigit():
                cell_a.value = no_counter
                no_counter += 1
    
    # Total formula
    lr = ws.max_row
    for ci in range(5, 11):
        col_ltr = get_column_letter(ci)
        ws.cell(lr, ci).value = f"=SUM({col_ltr}{data_start}:{col_ltr}{lr-1})"
        ws.cell(lr, ci).number_format = "#,##0"
    
    return safe


def _write_rekap_log(log_path: str, df: pd.DataFrame, config_map: Dict, not_found: Set) -> int:
    """Write log file untuk SUBCLASS tidak ditemukan."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    
    if not not_found:
        return 0
    
    H = [str(c).upper().strip() for c in df.columns]
    iSub = find_column_index(H, ["SUBCLASS", "SUBCLASS_CODE"])
    iOwn = find_column_index(H, ["OWNER_ASSET", "OWNER_ASET", "OWNER"])
    iSt = find_column_index(H, ["STATUS"])
    iEM = find_column_index(H, ["END_MEASURE", "PANJANG", "LENGTH"])
    iKap = find_column_index(H, ["KAPASITAS", "KAPASITAS_TRAFO", "RATED_CAPACITY"])
    
    rows = [
        row for _, row in df.iterrows()
        if iSub is not None and
           str(row.iloc[iSub] or "").strip().upper() in not_found
    ]
    
    if not rows:
        return 0
    
    wb = Workbook()
    ws = wb.active
    ws.title = "LOG_NOT_FOUND"
    
    hdrs = ["NO", "OWNER_ASSET", "SUBCLASS", "STATUS", "END_MEASURE", "KAPASITAS", "KETERANGAN"]
    hfill = PatternFill("solid", fgColor="CC0000")
    hfont = Font(bold=True, color="FFFFFF")
    
    for ci, h in enumerate(hdrs, 1):
        c = ws.cell(1, ci, h)
        c.fill = hfill
        c.font = hfont
        c.alignment = Alignment(horizontal="center")
    
    for ri, row in enumerate(rows, 2):
        ws.cell(ri, 1, ri - 1)
        ws.cell(ri, 2, row.iloc[iOwn] if iOwn is not None else "")
        ws.cell(ri, 3, row.iloc[iSub] if iSub is not None else "")
        ws.cell(ri, 4, row.iloc[iSt] if iSt is not None else "")
        ws.cell(ri, 5, row.iloc[iEM] if iEM is not None else "")
        ws.cell(ri, 6, row.iloc[iKap] if iKap is not None else "")
        ws.cell(ri, 7, "SUBCLASS TIDAK ADA DI CONFIG")
    
    for col in ws.columns:
        mw = 10
        for cell in col:
            if cell.value:
                mw = max(mw, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(mw + 4, 50)
    
    wb.save(log_path)
    return len(rows)


class RekapEngine(BaseEngine):
    """Engine untuk generate rekap report."""
    
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
    
    def run(
        self,
        template_path: str,
        master_path: str,
        uid: str,
        filter_mode: int,
        out_rekap: str,
        out_log: str,
        selected_sheets: Optional[List[str]] = None
    ) -> Dict[str, Any]:
        """
        Run rekap process.
        
        Returns: {owners, not_found, elapsed}
        """
        from openpyxl import load_workbook
        
        self._start_timer()
        
        try:
            self._flog(f"START Rekap | template={template_path} | master={master_path}")
            
            # Load config
            self._log_ui("Membaca template config...", "info")
            config_map, owner_map, template_rows, tmpl_name = _load_rekap_config(template_path)
            self._log_ui(
                f"  ✓ {len(config_map)} SUBCLASS mapping, {len(template_rows)} baris template",
                "success"
            )
            self._update_progress(10)
            
            # Read master data
            self._log_ui("Membaca file master data...", "info")
            xl = pd.ExcelFile(master_path)
            sheets_to_read = selected_sheets if selected_sheets else xl.sheet_names
            dfs = []
            for s in sheets_to_read:
                if s in xl.sheet_names:
                    dfs.append(pd.read_excel(master_path, sheet_name=s, dtype=str))
            
            if not dfs:
                raise ValueError("Tidak ada sheet yang bisa dibaca dari file master!")
            
            df = pd.concat(dfs, ignore_index=True) if len(dfs) > 1 else dfs[0]
            df.columns = [str(c).strip() for c in df.columns]
            self._log_ui(f"  ✓ {len(df):,} baris dibaca", "success")
            self._update_progress(25)
            
            # Process data
            self._log_ui("Memproses data...", "info")
            summary, not_found = _process_rekap_data(df, config_map, owner_map)
            if not summary:
                raise ValueError("Tidak ada data valid. Cek kolom SUBCLASS dan CONFIG sheet.")
            
            owners = sorted(summary.keys())
            self._log_ui(f"  ✓ {len(owners)} owner ditemukan", "success")
            self._update_progress(45)
            
            # Create output
            self._log_ui("Membuat file rekap Excel...", "info")
            wb_tpl = load_workbook(template_path)
            ws_tpl = wb_tpl[tmpl_name] if tmpl_name in wb_tpl.sheetnames else None
            if ws_tpl is None:
                raise ValueError(f"Sheet template '{tmpl_name}' tidak ditemukan!")
            
            from openpyxl import Workbook as WB
            wb_out = WB()
            wb_out.remove(wb_out.active)
            
            os.makedirs(Path(out_rekap).parent, exist_ok=True)
            
            for i, owner in enumerate(owners):
                sname = _write_rekap_sheet(
                    wb_out,
                    ws_tpl,
                    owner,
                    uid,
                    summary[owner],
                    template_rows,
                    filter_mode
                )
                self._log_ui(f"  ✓ Sheet: {sname}", "success")
                self._flog(f"  Sheet written: {sname}")
                self._update_progress(45 + int((i + 1) / len(owners) * 45))
            
            wb_out.save(out_rekap)
            self._log_ui("  ✓ File rekap disimpan", "success")
            
            # Write log
            n_log = 0
            if not_found:
                os.makedirs(Path(out_log).parent, exist_ok=True)
                n_log = _write_rekap_log(out_log, df, config_map, not_found)
                self._log_ui(
                    f"  ⚠ {n_log} baris SUBCLASS tidak ditemukan → {Path(out_log).name}",
                    "warning"
                )
                self._flog(f"Log written: {n_log} rows -> {out_log}", "warning")
            
            elapsed = self._get_elapsed()
            self._flog(f"DONE Rekap: owners={len(owners)} elapsed={elapsed}s")
            self._update_progress(100)
            
            return {
                "owners": len(owners),
                "not_found": n_log,
                "elapsed": elapsed
            }
        
        except Exception as e:
            self._flog(f"CRITICAL ERROR: {str(e)}", "error")
            raise